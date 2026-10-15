"""Conservative extraction only: no summarization, fact inference or OCR guesses."""
from hashlib import sha256
from html.parser import HTMLParser
from io import BytesIO
import json
import re
from zipfile import ZipFile,is_zipfile

from models.sources import ExtractedDocument,SourceSegment

MAX_FILE_BYTES=20*1024*1024
MAX_ARCHIVE_BYTES=80*1024*1024
MAX_PAGES=300


def clean_text(value):
    # Retain indentation, internal spacing, repeated labels and blank lines.
    return value.replace('\r\n','\n').replace('\r','\n')


def decode_text(value,media_type):
    warnings=[]
    charset=re.search(r'charset\s*=\s*["\']?([^;"\'\s]+)',media_type,re.I)
    encoding=charset.group(1) if charset else ('utf-16' if value.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig')
    try:return value.decode(encoding),warnings
    except (UnicodeDecodeError,LookupError):
        warnings.append('Text decoding used UTF-8 replacement characters; original bytes are retained.')
        return value.decode('utf-8',errors='replace'),warnings


class ConservativeHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts=[];self.tables=[];self.table_stack=[];self.links=[];self.warnings=[];self.images=False

    def handle_starttag(self,tag,attrs):
        attrs=dict(attrs)
        if tag in {'p','div','section','article','h1','h2','h3','h4','h5','h6','ul','ol','pre','blockquote','table','tr'}:self.parts.append('\n')
        if tag=='br':self.parts.append('\n')
        if tag=='li':self.parts.append('\n- ')
        if tag=='a':self.links.append(attrs.get('href'))
        if tag=='table':self.table_stack.append({'rows':[],'row':None,'cell':None,'spans':[]})
        if tag=='tr' and self.table_stack:self.table_stack[-1]['row']=[]
        if tag in {'td','th'} and self.table_stack:
            table=self.table_stack[-1]
            if table['cell'] is not None:
                if table['row'] is None:table['row']=[]
                table['row'].append(''.join(table['cell']))
                self.warnings.append('Unclosed HTML cell; observed text retained, layout needs verification.')
            self.table_stack[-1]['cell']=[]
            self.table_stack[-1]['spans'].append({'rowspan':attrs.get('rowspan','1'),'colspan':attrs.get('colspan','1'),'header':tag=='th'})
        if tag=='img':
            self.images=True
            self.parts.append('\n[Image: '+attrs.get('alt','')+'] '+attrs.get('src','')+'\n')

    def handle_data(self,data):
        self.parts.append(data)
        for table in self.table_stack:
            if table['cell'] is not None:table['cell'].append(data)

    def handle_endtag(self,tag):
        if tag=='a' and self.links:
            link=self.links.pop()
            if link:
                self.parts.append(' ['+link+']')
                for table in self.table_stack:
                    if table['cell'] is not None:table['cell'].append(' ['+link+']')
        if tag in {'td','th'}:
            self.parts.append('\t')
            if self.table_stack:
                table=self.table_stack[-1]
                if table['cell'] is not None:
                    if table['row'] is None:table['row']=[]
                    table['row'].append(''.join(table['cell']))
                    table['cell']=None
        if tag=='tr' and self.table_stack:
            table=self.table_stack[-1]
            if table['row'] is not None:table['rows'].append(table['row']);table['row']=None
        if tag=='table' and self.table_stack:
            table=self.table_stack.pop()
            if table['cell'] is not None:
                if table['row'] is None:table['row']=[]
                table['row'].append(''.join(table['cell']))
                self.warnings.append('Unclosed HTML cell; original markup retained.')
            if table['row'] is not None:table['rows'].append(table['row'])
            self.tables.append(table)
        if tag in {'p','div','section','article','h1','h2','h3','h4','h5','h6','li','pre','blockquote','table','tr'}:self.parts.append('\n')

    def result(self):
        if self.table_stack:
            self.warnings.append('Malformed/unclosed HTML table; original HTML retained.')
            self.tables.extend(self.table_stack)
        return ''.join(self.parts)


def table_text(cells):
    # Empty/None cells are represented explicitly in structured cells, not filled.
    return '\n'.join('\t'.join('' if cell is None else str(cell) for cell in row) for row in cells)


def extract_document(value,media_type):
    if len(value)>MAX_FILE_BYTES:raise ValueError('Document exceeds configured extraction size limit.')
    digest=sha256(value).hexdigest()
    mime=media_type.split(';',1)[0].lower().strip()
    warnings=[];segments=[];raw='';status='extracted'
    try:
        if mime in {'text/html','application/xhtml+xml'}:
            raw,warnings=decode_text(value,media_type)
            parser=ConservativeHTML();parser.feed(raw);parser.close();text=parser.result()
            warnings+=parser.warnings
            if warnings:status='partial'
            segments.append(SourceSegment(locator='html:body',original_text=text,cleaned_text=clean_text(text)))
            for index,table in enumerate(parser.tables,1):
                cells=table['rows'];text=table_text(cells)
                segments.append(SourceSegment(locator=f'html:table:{index}',kind='table',original_text=text,cleaned_text=clean_text(text),cells=cells,layout={'cell_spans':table['spans']},warnings=['Merged cells are recorded, not expanded or guessed.']))
            if parser.images:status='requires_multimodal';warnings.append('HTML image/diagram content needs visual inspection; image labels/URLs retained.')
        elif mime.startswith('text/') or mime=='application/json':
            raw,warnings=decode_text(value,media_type)
            segments.append(SourceSegment(locator='text:1',original_text=raw,cleaned_text=clean_text(raw)))
        elif mime=='application/pdf':
            import pdfplumber
            with pdfplumber.open(BytesIO(value)) as document:
                if len(document.pages)>MAX_PAGES:status='partial';warnings.append('Page extraction limit reached; complete original file retained.')
                for index,page in enumerate(document.pages[:MAX_PAGES],1):
                    text=page.extract_text(layout=True) or ''
                    segments.append(SourceSegment(locator=f'page:{index}',original_text=text,cleaned_text=clean_text(text),layout={'width':page.width,'height':page.height}))
                    try:
                        for table_index,cells in enumerate(page.extract_tables(),1):
                            text=table_text(cells)
                            segments.append(SourceSegment(locator=f'page:{index}:table:{table_index}',kind='table',original_text=text,cleaned_text=clean_text(text),cells=cells,warnings=['PDF cell grouping is extraction output, not verified examination facts.']))
                    except Exception:warnings.append(f'Page {index}: table extraction failed; page text/original retained.');status='partial'
                    if page.images or not page.chars:
                        warnings.append(f'Page {index}: images or absent text require multimodal inspection; no OCR inferred.')
                        if status!='partial':status='requires_multimodal'
                    if page.lines or page.curves:
                        warnings.append(f'Page {index}: vector layout/diagrams may require visual verification.')
                warnings.append('PDF reading order, borderless tables and merged cells may be ambiguous; verify against the original.')
            raw='\n\n'.join(s.original_text for s in segments if s.kind=='text')
        elif mime in {'application/vnd.openxmlformats-officedocument.wordprocessingml.document','application/vnd.openxmlformats-officedocument.presentationml.presentation'}:
            if not is_zipfile(BytesIO(value)):raise ValueError('Not an OOXML archive.')
            with ZipFile(BytesIO(value)) as archive:
                entries=archive.infolist()
                if len(entries)>10000 or sum(entry.file_size for entry in entries)>MAX_ARCHIVE_BYTES:raise ValueError('Archive extraction limit exceeded.')
            if mime.endswith('wordprocessingml.document'):
                from docx import Document
                from docx.table import Table
                document=Document(BytesIO(value))
                for index,item in enumerate(document.iter_inner_content(),1):
                    if isinstance(item,Table):
                        cells=[[cell.text for cell in row.cells] for row in item.rows];text=table_text(cells)
                        segments.append(SourceSegment(locator=f'section:block:{index}:table',kind='table',original_text=text,cleaned_text=clean_text(text),cells=cells,warnings=['Merged cells may repeat; nested table layout should be checked against the original.']))
                        def nested(tables,prefix):
                            for number,table in enumerate(tables,1):
                                cells=[[cell.text for cell in row.cells] for row in table.rows];text=table_text(cells)
                                segments.append(SourceSegment(locator=f'{prefix}:nested:{number}',kind='table',original_text=text,cleaned_text=clean_text(text),cells=cells))
                                for row_number,row in enumerate(table.rows):
                                    for cell_number,cell in enumerate(row.cells):nested(cell.tables,f'{prefix}:nested:{number}:row:{row_number}:cell:{cell_number}')
                        for row_number,row in enumerate(item.rows):
                            for cell_number,cell in enumerate(row.cells):nested(cell.tables,f'section:block:{index}:row:{row_number}:cell:{cell_number}')
                    else:
                        text=item.text
                        links=[node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id') for node in item._p.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink')]
                        urls=[document.part.rels[link].target_ref for link in links if link and link in document.part.rels]
                        segments.append(SourceSegment(locator=f'section:block:{index}',original_text=text,cleaned_text=clean_text(text)+'\n'.join(' ['+url+']' for url in urls),layout={'style':item.style.name,'links':urls}))
                for label,part in [(f'section:{number}:header',section.header) for number,section in enumerate(document.sections,1)]+[(f'section:{number}:footer',section.footer) for number,section in enumerate(document.sections,1)]:
                    for index,paragraph in enumerate(part.paragraphs,1):
                        if paragraph.text:segments.append(SourceSegment(locator=f'{label}:paragraph:{index}',original_text=paragraph.text,cleaned_text=clean_text(paragraph.text)))
                if document.inline_shapes:status='requires_multimodal';warnings.append('Word images/diagrams require visual inspection.')
                warnings.append('Word text boxes, footnotes and embedded objects may need original-file inspection.')
            else:
                from pptx import Presentation
                from pptx.enum.shapes import MSO_SHAPE_TYPE
                presentation=Presentation(BytesIO(value))
                def visit(shapes,prefix):
                    nonlocal status
                    for index,shape in enumerate(shapes,1):
                        locator=f'{prefix}:shape:{index}'
                        if shape.shape_type==MSO_SHAPE_TYPE.GROUP:visit(shape.shapes,locator);continue
                        if shape.has_text_frame:
                            text=shape.text_frame.text
                            urls=[run.hyperlink.address for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.hyperlink.address]
                            segments.append(SourceSegment(locator=locator,original_text=text,cleaned_text=clean_text(text)+''.join(' ['+url+']' for url in urls),layout={'left':shape.left,'top':shape.top,'links':urls}))
                        if shape.has_table:
                            cells=[[cell.text for cell in row.cells] for row in shape.table.rows];text=table_text(cells)
                            spans=[[{'merged_origin':cell.is_merge_origin,'spanned':cell.is_spanned,'rows':cell.span_height,'columns':cell.span_width} for cell in row.cells] for row in shape.table.rows]
                            segments.append(SourceSegment(locator=locator+':table',kind='table',original_text=text,cleaned_text=clean_text(text),cells=cells,layout={'cell_spans':spans}))
                        if shape.shape_type==MSO_SHAPE_TYPE.PICTURE or shape.has_chart:
                            status='requires_multimodal';warnings.append(locator+': image/chart needs visual inspection.')
                for slide_index,slide in enumerate(presentation.slides,1):
                    visit(slide.shapes,f'slide:{slide_index}')
                    if slide.has_notes_slide:
                        frame=slide.notes_slide.notes_text_frame
                        if frame and frame.text:segments.append(SourceSegment(locator=f'slide:{slide_index}:notes',kind='notes',original_text=frame.text,cleaned_text=clean_text(frame.text)))
                warnings.append('Slide shape order/embedded objects may need visual verification against the original.')
            raw='\n\n'.join(segment.original_text for segment in segments)
        else:
            status='unsupported';warnings.append('Unsupported format; original file/reference retained without inferred text.')
    except ImportError:
        status='unsupported';warnings.append('Optional extraction library unavailable; original file/reference retained.')
    except Exception:
        status='failed';warnings.append('Extraction failed; original file/reference retained. No source content discarded or facts inferred.')
    cleaned='\n\n'.join(segment.cleaned_text for segment in segments)
    return ExtractedDocument(media_type=media_type,status=status,original_text=raw,cleaned_text=cleaned,segments=segments,warnings=warnings,original_sha256=digest)
