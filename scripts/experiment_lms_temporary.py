"""Opt-in installed-tool sandbox probe. Synthetic PDF/API; no auth or live services."""
import argparse
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import io
from zipfile import ZipFile,ZIP_STORED

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'backend'))
from services.lms_temporary import private_scratch
from services.document_extraction import extract_document


def pdf_with_padding(padding):
    # Large unused PDF object mimics binary-heavy course files without OCR.
    content=b'BT /F1 12 Tf 72 720 Td (Queues process entries first in first out.) Tj ET'
    objects=[b'<< /Type /Catalog /Pages 2 0 R >>',
             b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
             b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>',
             b'<< /Length '+str(len(content)).encode()+b' >>\nstream\n'+content+b'\nendstream',
             b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
             b'<< /Length '+str(padding).encode()+b' >>\nstream\n'+b' '*padding+b'\nendstream']
    value=bytearray(b'%PDF-1.4\n');offsets=[0]
    for index,obj in enumerate(objects,1):
        offsets.append(len(value));value.extend(str(index).encode()+b' 0 obj\n'+obj+b'\nendobj\n')
    xref=len(value);value.extend(b'xref\n0 7\n0000000000 65535 f \n')
    for offset in offsets[1:]:value.extend(f'{offset:010d} 00000 n \n'.encode())
    value.extend(f'trailer\n<< /Size 7 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode())
    return bytes(value)


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--run',action='store_true')
    p.add_argument('--package-root',type=Path)
    args=p.parse_args()
    if not args.run:
        print('No package, settings, files or services accessed.');return 0
    if args.package_root is None:
        from services.lms_startup import configured_environment
        args.package_root=Path(configured_environment()['LMS_MCP_PACKAGE_ROOT'])
    package=args.package_root.resolve()
    manifest=json.loads((package/'package.json').read_text(encoding='utf-8'))
    if manifest.get('name')!='brightspace-mcp-server' or manifest.get('version')!='3.13.18':
        raise ValueError('The probe requires reviewed MCP version 3.13.18.')
    node=shutil.which('node.exe' if os.name=='nt' else 'node')
    if not node:raise ValueError('Node executable is required.')
    source_module=package/'build/tools/download-file.js'
    original_digest=hashlib.sha256(source_module.read_bytes()).hexdigest()
    reports=[]
    with tempfile.TemporaryDirectory(dir=ROOT/'tests',prefix='quiz-test-large-') as folder:
        base=Path(folder);scratch=base/'private';sentinel=base/'unrelated.txt';sentinel.write_text('untouched')
        for index,(mode,megabytes,header,expected,kind) in enumerate([
            ('inline',12,'accurate','refused','pdf'),
            ('bounded_disk',12,'accurate','extracted','pdf'),
            ('bounded_disk',12,'missing','extracted','pdf'),
            ('bounded_disk',21,'accurate','refused','pdf'),
            ('bounded_disk',21,'missing','refused','pdf'),
            ('bounded_disk',21,'understated','refused','pdf'),
            ('bounded_disk',12,'accurate','extracted','docx'),
            ('bounded_disk',12,'missing','extracted','pptx'),
        ],1):
            source=base/('synthetic.'+kind)
            if kind=='pdf':source.write_bytes(pdf_with_padding(megabytes*1024*1024))
            else:
                body=io.BytesIO()
                if kind=='docx':
                    from docx import Document
                    office=Document();office.add_paragraph('Queues process entries first in first out.')
                else:
                    from pptx import Presentation
                    from pptx.util import Inches
                    office=Presentation();slide=office.slides.add_slide(office.slide_layouts[6])
                    slide.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(1)).text='Queues process entries first in first out.'
                office.save(body)
                with ZipFile(body,'a') as archive:archive.writestr('synthetic-padding.bin',b' '*megabytes*1024*1024,compress_type=ZIP_STORED)
                source.write_bytes(body.getvalue())
            with private_scratch(scratch) as job:
                jobpath=job.directory
                # No bot/Google/IBM/school settings passed to the synthetic child.
                env={k:os.environ[k] for k in ('SystemRoot','WINDIR','COMSPEC','PATH') if k in os.environ}
                env.update(HOME=str(base),USERPROFILE=str(base),TEMP=str(base),TMP=str(base),
                           D2L_NO_UPDATE_CHECK='1',D2L_SESSION_DIR=str(base/'unused-synthetic-session'),
                           AI_STUDY_MCP_PACKAGE_ROOT=str(package))
                child=subprocess.run([node,'--import',(ROOT/'backend/services/lms_runtime/private_preload.mjs').as_uri(),
                                      str(ROOT/'scripts/lms_temporary_probe.mjs'),str(package),
                                      str(source),str(job.directory),mode,header],
                                     env=env,capture_output=True,text=True,encoding='utf-8',timeout=60)
                if child.returncode:
                    # This child receives synthetic data only; fixed error classes
                    # are enough to diagnose the harness without printing payloads.
                    import re
                    failures=re.findall(r'(?:Error|TypeError|SyntaxError): [^\r\n]+',child.stderr)
                    print('Synthetic child failure categories:',failures[-2:])
                    raise RuntimeError('Synthetic Node probe failed; payloads suppressed.')
                value=json.loads(child.stdout);result=value['result']
                if not value['preload']:raise AssertionError('Production preload was not installed.')
                failed=bool(result.get('isError'))
                extracted=False
                if expected=='extracted':
                    if failed:raise AssertionError('Expected synthetic download success.')
                    metadata=json.loads(result['content'][0]['text'])
                    if Path(metadata['filePath'])!=job.target:raise AssertionError('Unexpected output path.')
                    document=extract_document(job.read_bytes(metadata['mimeType']),metadata['mimeType'])
                    extracted='Queues process entries first in first out.' in document.cleaned_text
                    if not extracted:raise AssertionError('PDF text not extracted.')
                else:
                    if not failed or job.target.exists():raise AssertionError('Size refusal did not fail closed.')
                    if list(job.directory.iterdir()):raise AssertionError('Failed download left partial data.')
                if value['reads']!=1:raise AssertionError('Unexpected source request count.')
            if jobpath.exists():raise AssertionError('Private scratch job was not cleaned.')
            reports.append({'case':index,'mode':mode,'format':kind,'synthetic_size_mib':megabytes,
                            'content_length':header,'outcome':expected,'extracted':extracted,
                            'stream_cancelled':value['cancelled'],'job_removed':True})
        if sentinel.read_text()!='untouched' or list(scratch.iterdir()):
            raise AssertionError('Cleanup touched unrelated data or left a job.')
    if hashlib.sha256(source_module.read_bytes()).hexdigest()!=original_digest:
        raise AssertionError('Installed upstream source changed.')
    report={'synthetic_only':True,'actual_upstream_version':'3.13.18','max_download_mib':20,
            'upstream_source_unchanged':True,'no_real_network_or_auth':True,'cases':reports}
    artifact=ROOT/'tests/artifacts/lms-temporary-probe.json'
    artifact.parent.mkdir(parents=True,exist_ok=True)
    artifact.write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(report));return 0


if __name__=='__main__':raise SystemExit(main())
