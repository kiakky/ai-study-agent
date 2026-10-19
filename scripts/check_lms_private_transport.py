"""Opt-in real MCP transport with upstream tool but synthetic API/files/actors."""
import argparse,asyncio,io,json,os,shutil,sys,tempfile
from contextlib import asynccontextmanager,redirect_stdout
from pathlib import Path
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'backend'),str(ROOT/'tests'),str(ROOT/'scripts')]


async def verify(package,node):
    from mcp import ClientSession,StdioServerParameters
    from mcp.client.stdio import stdio_client
    from services.lms_policy import ReadOnlySession,ReadOnlyViolation
    from services.lms_mcp import MCPBrightspaceBackend
    from services.lms_payloads import attribute,MCPDownloadLimit
    from services.lms_retrieval import BrightspaceRetrieval
    from services.source_store import SourceStore,LMSAccessDenied
    from models.sources import SourceMetadata,SourceQuery
    from quiz_refinement_fixture import Fixture
    from experiment_lms_temporary import pdf_with_padding
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from zipfile import ZipFile,ZIP_STORED
    reports=[]
    with tempfile.TemporaryDirectory(dir=ROOT/'tests',prefix='quiz-test-wire-') as directory:
        base=Path(directory);scratch=base/'private'
        for kind,mb in [('pdf',12),('docx',12),('pptx',12),('pdf',21)]:
            source=base/('synthetic.'+kind)
            if kind=='pdf':source.write_bytes(pdf_with_padding(mb*1024*1024))
            else:
                body=io.BytesIO()
                if kind=='docx':office=Document();office.add_paragraph('Queues process entries first in first out.')
                else:
                    office=Presentation();slide=office.slides.add_slide(office.slide_layouts[6])
                    slide.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(1)).text='Queues process entries first in first out.'
                office.save(body)
                with ZipFile(body,'a') as archive:archive.writestr('synthetic-padding.bin',b' '*mb*1024*1024,compress_type=ZIP_STORED)
                source.write_bytes(body.getvalue())
            fixture=Fixture()
            with redirect_stdout(io.StringIO()):fixture.setUp()
            try:
                fixture.store.save_course({'course_id':'501','name':'Synthetic','activity':'active'})
                sid=fixture.store.index('501:file:7',SourceMetadata(course_id='501',resource_id='7',
                    kind='document',title='Synthetic '+kind,raw_metadata={'topicId':7,'topicType':'file'}),{'topicId':7})['source_id']
                env={k:os.environ[k] for k in ('SystemRoot','WINDIR','COMSPEC','PATH') if k in os.environ}
                env.update(HOME=str(base),USERPROFILE=str(base),TEMP=str(base),TMP=str(base),
                    D2L_NO_UPDATE_CHECK='1',D2L_SESSION_DIR=str(base/'unused-synthetic-session'),AI_STUDY_MCP_PACKAGE_ROOT=str(package))
                params=StdioServerParameters(command=node,args=['--import',
                    (ROOT/'backend/services/lms_runtime/private_preload.mjs').as_uri(),
                    str(ROOT/'scripts/lms_temporary_probe.mjs'),str(package),str(source),str(scratch),'stdio','missing'],env=env)
                with open(os.devnull,'w') as errors:
                    async with stdio_client(params,errlog=errors) as (read,write):
                        async with ClientSession(read,write) as session:
                            init=await session.initialize()
                            info=attribute(init,'serverInfo','server_info')
                            guarded=ReadOnlySession(session,attribute(info,'version'),private_downloads=True)
                            tools=await guarded.list_tools()
                            # The public read boundary still denies caller-selected paths.
                            try:await guarded.call_tool('download_file',{'courseId':501,'topicId':7,'downloadPath':str(base)})
                            except ReadOnlyViolation:pass
                            else:raise AssertionError('Public path argument was accepted.')
                            class OpenBackend(MCPBrightspaceBackend):
                                @asynccontextmanager
                                async def batch(self):yield self
                            backend=OpenBackend(allow_document_downloads=True);backend.session=guarded
                            backend.schemas={t.name:attribute(t,'inputSchema','input_schema',default={}) for t in tools.tools}
                            retrieval=BrightspaceRetrieval(101,backend,fixture.store)
                            def limited(error):return isinstance(error,MCPDownloadLimit) or any(limited(e) for e in getattr(error,'exceptions',()))
                            with patch('services.lms_temporary.SCRATCH_ROOT',scratch):
                                try:result=await retrieval.retrieve_course_document(SourceQuery(source_id=sid))
                                except Exception as error:
                                    if mb!=21 or not limited(error):raise
                                    result=None
                            if mb==12:
                                text='\n'.join(s['cleaned_text'] for s in result['segments'])
                                assert 'Queues process entries first in first out.' in text
                                assert not result['original_file_available']
                                assert fixture.store.cached_document(sid)['original_file_reference'] is None
                                # Cached text remains usable without another transport request.
                                again=await retrieval.retrieve_course_document(SourceQuery(source_id=sid))
                                assert again['cache_used']
                            else:assert fixture.store.cached_document(sid) is None
                            other=SourceStore(202,owner_id=101,cache_root=fixture.store.cache_root)
                            try:other.evidence(sid)
                            except LMSAccessDenied:pass
                            else:raise AssertionError('Other actor obtained owner evidence.')
                assert not scratch.exists() or not list(scratch.iterdir())
                reports.append({'format':kind,'mib':mb,'text_cached_without_binary':mb==12,
                    'limit_refused':mb==21,'job_removed':True,'other_actor_denied':True,'public_paths_denied':True})
            finally:fixture.doCleanups()
    return reports


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--run',action='store_true')
    parser.add_argument('--package-root',type=Path);args=parser.parse_args()
    if not args.run:print('No settings, package, database or services accessed.');return 0
    if args.package_root is None:
        from services.lms_startup import configured_environment
        args.package_root=Path(configured_environment()['LMS_MCP_PACKAGE_ROOT'])
    node=shutil.which('node.exe' if os.name=='nt' else 'node')
    rows=asyncio.run(verify(args.package_root.resolve(),node))
    report={'synthetic_only':True,'real_mcp_transport':True,'cases':rows}
    (ROOT/'tests/artifacts/lms-private-transport.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(report));return 0


if __name__=='__main__':raise SystemExit(main())
