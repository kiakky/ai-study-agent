import base64
from contextlib import asynccontextmanager,closing
from copy import deepcopy
from datetime import datetime,timezone
from io import BytesIO
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch,AsyncMock

from test_api import api,database,signed_data,TEST_TOKEN
from fastapi.testclient import TestClient
from models.sources import CourseQuery,SourceQuery,SearchQuery,SourceMetadata
from models.exams import ExamInformation
from services.document_extraction import extract_document,clean_text
from services.lms_mcp import MCPBrightspaceBackend,MCPUnavailable,response_payload,document_bytes
from services.lms_retrieval import BrightspaceRetrieval
from services.source_store import SourceStore,LMSAccessDenied
from services.brightspace import server_parameters
from routes import lms as lms_routes

ROOT=Path(__file__).resolve().parents[1]
DOCX='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
PPTX='application/vnd.openxmlformats-officedocument.presentationml.presentation'
HTML='''<h1>Assessment requirements</h1><p>Approximate exam: Weeks 14–17. Exact date: TBC.</p>
<table><tr><th>Format</th><th>Questions</th><th>Marks</th></tr>
<tr><td>MCQ</td><td>40</td><td>40</td></tr><tr><td>Short answer</td><td>10</td><td>30</td></tr>
<tr><td>Practical</td><td>10</td><td></td></tr></table>
<p>Instructions unchanged.</p><p>Instructions unchanged.</p><ul><li>Chapter 4: SQL JOINs</li></ul>
<a href="https://reference.example/lecture">Reference</a><pre>    SELECT *\n    FROM samples;</pre>'''


def pdf_fixture(image_only=False):
    """Small valid synthetic PDF, including a bordered assessment table or image."""
    content=b'q 100 0 0 100 40 600 cm /Im0 Do Q' if image_only else b'BT /F1 12 Tf 40 760 Td (Exam period: Weeks 14-17; confirmed date: TBC) Tj ET\n'
    if not image_only:
        for row,values in enumerate([['Format','Questions','Marks'],['MCQ','40','40'],['Short answer','10','30']]):
            for column,value in enumerate(values):content+=f'BT /F1 12 Tf {45+column*130} {715-row*25} Td ({value}) Tj ET\n'.encode()
        for x in [40,170,300,430]:content+=f'{x} 650 m {x} 735 l S\n'.encode()
        for y in [650,675,700,735]:content+=f'40 {y} m 430 {y} l S\n'.encode()
    objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
             b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> /XObject << /Im0 6 0 R >> >> /Contents 5 0 R >>',
             b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
             b'<< /Length '+str(len(content)).encode()+b' >>\nstream\n'+content+b'\nendstream',
             b'<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Length 3 >>\nstream\n\x00\x00\x00\nendstream']
    data=b'%PDF-1.4\n';offsets=[0]
    for index,obj in enumerate(objects,1):offsets.append(len(data));data+=f'{index} 0 obj\n'.encode()+obj+b'\nendobj\n'
    xref=len(data);data+=f'xref\n0 {len(offsets)}\n'.encode()+b'0000000000 65535 f \n'
    for offset in offsets[1:]:data+=f'{offset:010d} 00000 n \n'.encode()
    return data+f'trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode()


class ExtractionTests(unittest.TestCase):
    def test_html_keeps_values_blank_cells_urls_repetition_and_code(self):
        result=extract_document(HTML.encode(),'text/html')
        self.assertEqual(result.original_text,HTML)
        self.assertIn('Weeks 14–17',result.cleaned_text)
        self.assertIn('Exact date: TBC',result.cleaned_text)
        table=next(s for s in result.segments if s.kind=='table')
        self.assertEqual(table.cells[1],['MCQ','40','40'])
        self.assertEqual(table.cells[3],['Practical','10',''])
        self.assertGreaterEqual(result.cleaned_text.count('Instructions unchanged.'),2)
        self.assertIn('https://reference.example/lecture',result.cleaned_text)
        self.assertIn('    SELECT *\n    FROM samples;',result.cleaned_text)
        self.assertEqual(clean_text('a\r\n\r\n    b  c'),'a\n\n    b  c')

    def test_malformed_html_images_and_decoding_are_flagged_not_guessed(self):
        malformed=extract_document(b'<table><tr><td>MCQ<td>40</table>','text/html')
        self.assertIn('MCQ',malformed.cleaned_text);self.assertIn('40',malformed.cleaned_text)
        self.assertEqual(malformed.status,'partial')
        result=extract_document(b'<img src="table.png" alt="Exam table">','text/html')
        self.assertEqual(result.status,'requires_multimodal')
        self.assertIn('table.png',result.cleaned_text)
        result=extract_document(b'Weight 40% \xff date TBC','text/plain')
        self.assertTrue(result.warnings);self.assertIn('40%',result.cleaned_text)

    def test_pdf_tables_and_image_pages_preserve_uncertainty(self):
        result=extract_document(pdf_fixture(),'application/pdf')
        self.assertIn('Weeks 14-17',result.original_text)
        tables=[s for s in result.segments if s.kind=='table']
        self.assertTrue(tables)
        self.assertTrue(any('MCQ' in cell for row in tables[0].cells for cell in row if cell))
        self.assertTrue(result.warnings)
        image=extract_document(pdf_fixture(True),'application/pdf')
        self.assertEqual(image.status,'requires_multimodal')
        self.assertEqual(image.original_text,'')
        self.assertFalse(any(s.kind=='table' for s in image.segments))

    def test_docx_nested_tables_headings_and_pptx_tables_notes(self):
        from docx import Document
        document=Document();document.add_heading('Examination scope',1);document.add_paragraph('Exam week is approximate: 14–17. Date not confirmed.')
        table=document.add_table(rows=2,cols=2);table.cell(0,0).text='Type';table.cell(0,1).text='Marks';table.cell(1,0).text='MCQ';table.cell(1,1).text='40'
        nested=table.cell(1,0).add_table(rows=1,cols=2);nested.cell(0,0).text='Learning outcome';nested.cell(0,1).text='LO4'
        stream=BytesIO();document.save(stream)
        result=extract_document(stream.getvalue(),DOCX)
        self.assertIn('14–17',result.original_text);self.assertIn('LO4',result.original_text)
        self.assertTrue(any(s.layout.get('style')=='Heading 1' for s in result.segments))
        from pptx import Presentation
        from pptx.util import Inches
        presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text='Machine learning — Week 4 (not necessarily an exam course)'
        table=slide.shapes.add_table(2,2,Inches(1),Inches(2),Inches(5),Inches(2)).table
        table.cell(0,0).text='Questions';table.cell(0,1).text='Marks';table.cell(1,0).text='40';table.cell(1,1).text=''
        slide.notes_slide.notes_text_frame.text='Important source note, not a generated explanation.'
        stream=BytesIO();presentation.save(stream)
        result=extract_document(stream.getvalue(),PPTX)
        self.assertIn('Machine learning',result.original_text)
        self.assertTrue(any(s.kind=='notes' for s in result.segments))
        self.assertEqual(next(s for s in result.segments if s.kind=='table').cells[1],['40',''])

    def test_unsupported_and_corrupt_formats_keep_warning_hash(self):
        result=extract_document(b'opaque binary','application/octet-stream')
        self.assertEqual(result.status,'unsupported');self.assertTrue(result.original_sha256)
        result=extract_document(b'not pdf','application/pdf')
        self.assertEqual(result.status,'failed');self.assertTrue(result.warnings)


class FakeBackend:
    def __init__(self):
        self.calls=[];self.fail=set();self.documents={};self.catalog={'courses':[{'id':'old','name':'Legacy projects','isActive':False},{'id':'ml','name':'Machine learning support','isActive':True}]}
        self.content={'modules':[{'id':'week4','title':'Week 4','topics':[{'id':'lecture','title':'Lecture — ML models','url':'https://school.example/lecture.pptx','lastModified':'2026-10-12T09:00:00Z'}]}, {'id':'exam-module','title':'Assessments','topics':[{'id':'scope','title':'Assessment scope','html':HTML,'mimeType':'text/html'}]}]}
        self.announcements={'announcements':[{'id':'a1','title':'Assessment timing','body':'Possible exam in Weeks 14–17; date TBC.','modifiedAt':'2026-10-12T09:00:00Z'}]}

    @asynccontextmanager
    async def batch(self):yield self

    async def request(self,name,**arguments):
        self.calls.append((name,arguments))
        if name in self.fail:raise RuntimeError('Synthetic tool failure with private details not exposed')
        if name=='get_my_courses':payload=self.catalog
        elif name=='get_course_content':payload=self.content
        elif name=='get_announcements':payload=self.announcements
        elif name=='get_syllabus':payload={'text':'Course assistance for projects; no confirmed exam information.'}
        elif name=='get_assignments':payload={'assignments':[{'id':'p1','title':'Project','text':'Project weight 30%. Do not treat as an exam.'}]}
        elif name=='get_upcoming_due_dates':payload={'items':[]}
        elif name=='get_calendar_events':payload={'items':[{'title':'Assessment period','notes':'Dates not confirmed.'}]}
        elif name=='search_course':payload={'results':[{'id':'lecture','title':'Broader ML lecture match','query':arguments['query']}]}
        elif name=='download_file':payload={'base64':base64.b64encode(self.documents.get(arguments['resource']['resource_id'],pdf_fixture())).decode(),'mimeType':'application/pdf'}
        else:raise RuntimeError('Unsupported synthetic operation')
        return {'structuredContent':payload,'isError':False}


class RetrievalTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.folder=tempfile.TemporaryDirectory(prefix='quiz-test-',dir=ROOT/'tests');self.addCleanup(self.folder.cleanup)
        db=patch.object(database,'DB_PATH',Path(self.folder.name)/'sources.db');db.start();self.addCleanup(db.stop)
        database.initialize_database(backfill_approvals=False,include_lms=True)
        self.store=SourceStore(101,owner_id=101,cache_root=Path(self.folder.name)/'private')
        self.store.bind_authorized_developer_session('synthetic-authorized-school')
        self.backend=FakeBackend();self.service=BrightspaceRetrieval(101,self.backend,self.store)

    async def seed(self):return await self.service.invoke('initial_discovery',{})

    async def test_all_courses_and_module_relationships_without_downloads_or_exam_facts(self):
        result=await self.seed()
        self.assertEqual(len(result['catalog']['courses']),2)
        self.assertEqual(result['catalog']['courses'][0]['activity'],'active')
        self.assertFalse(any(name=='download_file' for name,args in self.backend.calls))
        resources=self.store.sources('old')
        self.assertEqual(next(s for s in resources if s['resource_id']=='lecture')['parent_id'],'week4')
        self.assertTrue(any(s['kind']=='assessments:item' for s in resources))
        self.assertEqual(result['courses'][0]['assessment_evidence']['exam_classification'],'not_interpreted')
        exam=ExamInformation(course_id='old',course_name='Legacy projects')
        self.assertEqual(exam.final_exam_exists.status,'missing');self.assertIsNone(exam.confirmed_start.value)

    async def test_pagination_catalog_and_unknown_shapes_preserve_source(self):
        original=self.backend.request
        async def paginated(name,**arguments):
            if name=='get_my_courses':
                return {'structuredContent':{'courses':[{'id':'two','name':'Older course'}] if arguments.get('cursor') else [{'id':'one','name':'Current course'}],'nextPageToken':None if arguments.get('cursor') else 'next'}}
            return await original(name,**arguments)
        self.backend.request=paginated
        result=await self.service.list_courses();self.assertEqual(len(result['courses']),2);self.assertTrue(result['enumeration_complete'])
        self.backend.request=lambda *args,**kwargs:None  # Replace below with async unknown response.
        async def unknown(*args,**kwargs):return {'content':[{'type':'text','text':'Unknown vendor response; preserve this exact source.'}]}
        self.backend.request=unknown
        result=await self.service.list_courses()
        self.assertTrue(result['warnings']);self.assertEqual(len(self.store.courses()),2)

    async def test_owner_gate_prevents_cache_and_live_calls_for_other_users(self):
        await self.seed();source=self.store.sources('ml')[0]['source_id']
        other=SourceStore(202,owner_id=101,cache_root=self.store.cache_root)
        service=BrightspaceRetrieval(202,self.backend,other)
        before=len(self.backend.calls)
        self.assertRaises(LMSAccessDenied,other.evidence,source)
        with self.assertRaises(LMSAccessDenied):await service.invoke('list_courses',{})
        self.assertEqual(len(self.backend.calls),before)
        with closing(database.get_connection()) as db,db:db.execute("UPDATE lms_connections SET status='disconnected'")
        with self.assertRaises(LMSAccessDenied):await self.service.invoke('retrieve_source_evidence',{'source_id':source})

    async def test_on_demand_full_document_cache_exact_duplicates_and_versioned_provenance(self):
        await self.seed();source=next(s for s in self.store.sources('ml') if s['resource_id']=='lecture')
        result=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id']))
        self.assertEqual(result['extraction_status'],'extracted')
        self.assertTrue(result['original_file_available'])
        revision=result['evidence']['revision_id']
        self.assertEqual(self.store.original_bytes(source['source_id'],revision),pdf_fixture())
        calls=len(self.backend.calls)
        cached=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id']))
        self.assertTrue(cached['cache_used']);self.assertEqual(len(self.backend.calls),calls)
        self.backend.content['modules'][0]['topics'][0]['lastModified']='2026-10-13T09:00:00Z'
        self.backend.content['modules'][0]['topics'][0]['title']='Updated lecture'
        await self.service.list_course_resources(CourseQuery(course_id='ml'))
        changed=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id']))
        self.assertFalse(changed['cache_used'])
        old=self.store.evidence(source['source_id'],revision)
        self.assertEqual(old['metadata']['title'],'Lecture — ML models')
        self.assertEqual(len(list((self.store.cache_root/self.store.connection()).glob('*.bin'))),1)

    async def test_changes_and_failed_lookups_never_delete_evidence(self):
        await self.seed();before=self.store.sources('ml')
        self.backend.fail={'get_course_content','get_announcements'}
        result=await self.service.check_course_updates(CourseQuery(course_id='ml'))
        self.assertTrue(result['warnings']);self.assertFalse(result['removals_inferred'])
        self.assertEqual(len(self.store.sources('ml')),len(before))
        self.backend.fail=set();self.backend.content={'modules':[]}
        result=await self.service.check_course_updates(CourseQuery(course_id='ml'))
        self.assertFalse(result['removals_inferred']);self.assertTrue(self.store.sources('ml'))

    async def test_html_table_evidence_and_search_not_exact_topic_matching(self):
        await self.seed();source=next(s for s in self.store.sources('ml') if s['resource_id']=='scope')
        document=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id']))
        evidence=self.store.evidence(source['source_id'],document['evidence']['revision_id'],'html:table:1')
        self.assertEqual(evidence['segments'][0]['cells'][1],['MCQ','40','40'])
        self.assertIsNone(evidence['evidence']['published_at'])
        result=await self.service.search_learning_materials(SearchQuery(course_id='ml',query='Week 4 models'))
        self.assertTrue(result['local_candidates']);self.assertIsNotNone(result['remote_results'])
        with self.assertRaises(LookupError):await self.service.invoke('list_course_resources',{'course_id':'not-authorized'})

    async def test_announcement_refresh_and_source_revisions_keep_conflicting_originals(self):
        await self.seed()
        source=next(s for s in self.store.sources('ml') if s['kind']=='announcements:item')
        first=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id']))
        self.assertIn('Weeks 14–17',first['original_text'])
        self.backend.announcements['announcements'][0]['body']='New notice: expected Weeks 12–13. Confirmed date still missing.'
        self.backend.announcements['announcements'][0]['modifiedAt']='2026-10-13T10:00:00Z'
        result=await self.service.check_course_updates(CourseQuery(course_id='ml'))
        self.assertTrue(any(item['change']=='changed' for item in result['assessment_sources']['source_references']))
        second=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id']))
        self.assertIn('Weeks 12–13',second['original_text'])
        old=self.store.evidence(source['source_id'],first['evidence']['revision_id'])
        self.assertIn('Weeks 14–17',old['original_text'])
        self.assertNotEqual(first['evidence']['revision_id'],second['evidence']['revision_id'])
        self.assertFalse(any(name=='download_file' for name,args in self.backend.calls))

    async def test_payload_change_without_timestamp_invalidates_cached_extraction(self):
        await self.seed();source=next(s for s in self.store.sources('ml') if s['kind']=='announcements:item')
        first=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id']))
        self.backend.announcements['announcements'][0]['body']='Updated approximate period, no confirmed date.'
        await self.service.check_course_updates(CourseQuery(course_id='ml'))
        second=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id']))
        self.assertIn('Updated approximate period',second['original_text'])
        self.assertFalse(second['cache_used'])
        self.assertNotEqual(first['evidence']['revision_id'],second['evidence']['revision_id'])
        aggregate=next(s for s in self.store.sources('ml') if s['kind']=='announcements')
        row=self.store.source(aggregate['source_id'])
        from models.sources import SourceMetadata
        metadata=SourceMetadata.model_validate(json.loads(row['metadata_json']))
        changed=self.store.index('ml:announcements',metadata,{'announcements':[],'notice':'changed'})
        self.assertEqual(changed['change'],'changed')
        unchanged=self.store.index('ml:announcements',metadata,{'announcements':[],'notice':'changed'})
        self.assertEqual(unchanged['change'],'unchanged')

    async def test_reverted_payload_is_a_new_revision_not_stale_cached_content(self):
        await self.seed();source=next(s for s in self.store.sources('ml') if s['kind']=='announcements:item')
        original=deepcopy(self.backend.announcements)
        first=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id']))
        self.backend.announcements['announcements'][0]['body']='Temporarily changed notice.'
        await self.service.check_course_updates(CourseQuery(course_id='ml'))
        changed=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id']))
        self.backend.announcements=original
        await self.service.check_course_updates(CourseQuery(course_id='ml'))
        reverted=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id']))
        self.assertEqual(reverted['original_text'],first['original_text'])
        self.assertNotEqual(reverted['evidence']['revision_id'],changed['evidence']['revision_id'])
        self.assertNotEqual(reverted['evidence']['revision_id'],first['evidence']['revision_id'])
        self.assertFalse(reverted['cache_used'])

    async def test_metadata_only_syllabus_is_retrievable_as_full_source(self):
        await self.seed();source=next(s for s in self.store.sources('ml') if s['kind']=='syllabus')
        result=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id']))
        self.assertIn('no confirmed exam information',result['original_text'])
        self.assertTrue(result['original_file_available'])

    async def test_corrupt_cached_original_is_preserved_and_recovery_is_non_destructive(self):
        await self.seed();source=next(s for s in self.store.sources('ml') if s['resource_id']=='lecture')
        first=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id']))
        revision=self.store.latest(source['source_id'],first['evidence']['revision_id'])
        original=self.store.cache_root/revision['original_file_reference']
        original.write_bytes(b'synthetic interrupted copy')
        self.assertRaises(LMSAccessDenied,self.store.original_bytes,source['source_id'],revision['id'])
        second=await self.service.retrieve_course_document(SourceQuery(source_id=source['source_id'],refresh=True))
        self.assertEqual(original.read_bytes(),b'synthetic interrupted copy')
        self.assertEqual(self.store.original_bytes(source['source_id'],second['evidence']['revision_id']),pdf_fixture())


class MCPAdapterTests(unittest.IsolatedAsyncioTestCase):
    async def test_mcp_child_does_not_inherit_application_secrets(self):
        import test_lms_upstream as upstream_fixtures
        with upstream_fixtures.fake_package() as package, patch.dict('os.environ',{'TELEGRAM_BOT_TOKEN':'synthetic','GOOGLE_API_KEY':'synthetic','OPENAI_API_KEY':'synthetic','D2L_HEADLESS':'true',
            'LMS_MCP_PACKAGE_ROOT':str(package),'LMS_SESSION_DIR':str(ROOT/'synthetic-session'),
            'LMS_NODE_EXECUTABLE':str(ROOT/'node.exe'),'LMS_TENANT_ORIGIN':'https://school.example','LMS_EXPECTED_ACCOUNT':'synthetic-owner'}):
            parameters=server_parameters()
        self.assertNotIn('TELEGRAM_BOT_TOKEN',parameters.env)
        self.assertNotIn('GOOGLE_API_KEY',parameters.env)
        self.assertNotIn('OPENAI_API_KEY',parameters.env)
        self.assertEqual(parameters.env['D2L_HEADLESS'],'false')
        self.assertEqual(parameters.env['D2L_NO_UPDATE_CHECK'],'1')
    async def test_schema_mapping_read_whitelist_and_unknown_required_fields(self):
        backend=MCPBrightspaceBackend()
        backend.session=SimpleNamespace(call_tool=AsyncMock(return_value={'structuredContent':{'courses':[]}}))
        backend.schemas={'get_course_content':{'type':'object','properties':{'courseId':{'type':'integer'},'maxDepth':{'type':'integer','maximum':7}},'required':['courseId']}}
        await backend.request('get_course_content',course_id='123')
        self.assertEqual(backend.session.call_tool.call_args.kwargs['arguments'],{'courseId':123,'maxDepth':7})
        with self.assertRaises(MCPUnavailable):await backend.request('delete_file',course_id='123')
        backend.schemas['get_course_content']['required'].append('unknown_required')
        with self.assertRaises(MCPUnavailable):await backend.request('get_course_content',course_id='123')

    async def test_download_paths_and_unapproved_hosts_are_not_opened(self):
        self.assertRaises(MCPUnavailable,document_bytes,{'structuredContent':{'path':'backend/credentials/private.json'}})
        backend=MCPBrightspaceBackend();backend.session=SimpleNamespace(call_tool=AsyncMock())
        backend.schemas={'download_file':{'type':'object','properties':{'url':{'type':'string'}},'required':['url']}}
        with self.assertRaises(MCPUnavailable):await backend.request('download_file',course_id='x',resource={'resource_id':'r','source_url':'https://unapproved.example/file.pdf'})
        self.assertFalse(backend.session.call_tool.called)
        self.assertRaises(MCPUnavailable,response_payload,{'isError':True,'content':[{'type':'text','text':'private details'}]})

    async def test_explicit_download_schema_does_not_confuse_topic_id_with_file_id(self):
        backend=MCPBrightspaceBackend(allow_document_downloads=True)
        backend.session=SimpleNamespace(call_tool=AsyncMock(return_value={'structuredContent':{'text':'synthetic doc'}}))
        backend.schemas={'download_file':{'type':'object','properties':{'fileId':{'type':'integer'}},'required':['fileId']}}
        with self.assertRaises(MCPUnavailable):await backend.request('download_file',course_id='123',resource={'resource_id':'456','raw_metadata':{'topicId':456}})
        self.assertFalse(backend.session.call_tool.called)
        with self.assertRaises(MCPUnavailable):await backend.request('download_file',course_id='123',resource={'resource_id':'456','raw_metadata':{'fileId':789}})
        self.assertFalse(backend.session.call_tool.called)
        backend.schemas={'download_file':{'type':'object','properties':{'courseId':{'type':'integer'},'topicId':{'type':'integer'}},'required':['courseId','topicId']}}
        await backend.request('download_file',course_id='123',resource={'raw_metadata':{'topicId':456}})
        self.assertEqual(backend.session.call_tool.call_args.kwargs['arguments'],{'courseId':123,'topicId':456})


class LMSMigrationTests(unittest.TestCase):
    def test_default_startup_does_not_deploy_lms_tables(self):
        with tempfile.TemporaryDirectory(prefix='quiz-test-',dir=ROOT/'tests') as folder:
            with patch.object(database,'DB_PATH',Path(folder)/'startup.db'):
                database.initialize_database(backfill_approvals=False)
                with closing(database.get_connection()) as db:
                    names={row[0] for row in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
                    self.assertNotIn('lms_connections',names)
                    self.assertIn('calendar_write_actions',names)

    def test_additive_schema_preserves_existing_synthetic_data_and_is_repeatable(self):
        import sqlite3
        with sqlite3.connect(':memory:') as db:
            db.executescript(database.SCHEMA.replace(database.LMS_SCHEMA,''))
            db.execute('INSERT INTO users (telegram_user_id) VALUES (?)',(101,))
            before=list(db.execute('SELECT * FROM users'))
            schemas=dict(db.execute("SELECT name,sql FROM sqlite_master WHERE type='table'"))
            db.executescript(database.LMS_SCHEMA)
            db.executescript(database.LMS_SCHEMA)
            after=dict(db.execute("SELECT name,sql FROM sqlite_master WHERE type='table'"))
            self.assertEqual(set(after)-set(schemas),{'lms_connections','lms_courses','lms_sources','lms_source_revisions'})
            self.assertTrue(all(after[name]==sql for name,sql in schemas.items()))
            self.assertEqual(list(db.execute('SELECT * FROM users')),before)
            self.assertEqual(db.execute('PRAGMA integrity_check').fetchone()[0],'ok')


class LMSAPITests(unittest.TestCase):
    def setUp(self):
        self.folder=tempfile.TemporaryDirectory(prefix='quiz-test-',dir=ROOT/'tests');self.addCleanup(self.folder.cleanup)
        db=patch.object(database,'DB_PATH',Path(self.folder.name)/'api.db');db.start();self.addCleanup(db.stop)
        token=patch.object(api,'BOT_TOKEN',TEST_TOKEN);token.start();self.addCleanup(token.stop)
        self.client=TestClient(api.app,base_url='http://localhost');self.client.__enter__();self.addCleanup(self.client.__exit__,None,None,None)

    def test_routes_keep_demo_users_independent_and_do_not_expose_owner_lms(self):
        self.assertEqual(self.client.get('/api/lms/status').status_code,401)
        self.assertEqual(self.client.post('/api/lms/tools/list_courses',json={}).status_code,401)
        with patch.dict('os.environ',{'TELEGRAM_OWNER_ID':'101'}):
            headers={'X-Telegram-Init-Data':signed_data(202)}
            self.assertFalse(self.client.get('/api/lms/status',headers=headers).json()['available'])
            self.assertEqual(self.client.post('/api/lms/tools/list_courses',headers=headers,json={}).status_code,403)
            self.assertEqual(self.client.get('/api/quizzes',headers=headers).status_code,200)


if __name__=='__main__':unittest.main()
